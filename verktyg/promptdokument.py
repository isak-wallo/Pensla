# Skriver bildpromptarna i GPT-stilen (se CLAUDE.md, "GPT-stilen") till två
# dokument att ge en bildgenerator:
#   promptar/bilder.md    för en agent som sparar filerna själv (Gemini i
#                         Antigravity): den sparar bilderna i inkorg/
#   promptar/bilder.docx  för en chatt (Copilot, ChatGPT) där man laddar ner
#                         bilderna själv
# Lägg till nya bilder i BILDER; filnamnet ska vara ledigt i bilder/farg/.
# Kör: python verktyg/promptdokument.py (kräver python-docx).
from docx import Document
from docx.shared import Pt, Cm
import os
ROT = os.path.join(os.path.dirname(os.path.abspath(__file__)), '..')
UT = os.path.join(ROT, 'promptar') + os.sep

# Vattnet, sanden, snön osv. blir annars brus eller försvinner vid spårningen
VATTEN = ("The water is one large calm area in soft blue with a few bold outlined wave lines. "
          "No reflections, no foam, no splashes, no sparkles.")
SAND = ("Sand and gravel are large smooth shapes in soft sandy beige, no single stones, no pebbles, "
        "no grainy texture, no flying dust.")
GATA = ("The street is one plain area with a few wide outlined markings, no paving tiles, "
        "no cobblestones, no grid of lines on the ground.")
GRUSVAG = ("The dirt road is one plain area with no stones, no ruts and no tire tracks. "
           "The fields are large simple areas, no rows of crops.")
BYGGE = ("The ground is one plain area of soft sandy beige, no tire tracks, no stones, no paving "
         "tiles. " + SAND)

# (rubrik, filnamn, motiv, proportioner, hela (det som ska synas), fordonet,
#  bakgrundens konturer, ränder/mönster, extra regler, mörka delar, detaljer,
#  bakgrund)
BILDER = [
    ("Grävmaskin", "gravare6",
     "a big yellow excavator on tracks at a construction site, digging into a sand pile with its bucket, seen from the side at a slight three-quarter angle.",
     "an excavator with a cab with big windows, a long bent arm with a bucket, wide tracks with a few large wheels inside",
     "the whole excavator with its arm and bucket", "vehicle",
     "the sand pile, the fence, traffic cones, the ground",
     "Stripes and warning patterns on the excavator",
     [BYGGE, "The tracks are one smooth band around a few large round wheels, no small tread pattern."],
     "the tracks, the arm's dark parts and the hydraulic cylinders",
     "an orange beacon on the cab roof, a side mirror, headlights, a step",
     "a simple half-built wall, a simple fence, two or three traffic cones, soft hills, a big rounded bush"),
    ("Dumper", "dumper5",
     "a big dump truck at a construction site, carrying a load of sand in its raised-edge dump body, seen from the side at a slight three-quarter angle.",
     "a dump truck with a cab with a big windshield and side window, a big box-shaped dump body, three big wheels on each side",
     "the whole truck", "vehicle",
     "the sand pile, the fence, traffic cones, the ground",
     "Stripes and warning patterns on the truck",
     [BYGGE],
     "tires and dark parts",
     "a beacon on the cab roof, side mirrors, headlights, steps up to the cab, a colored stripe along the cab",
     "a simple half-built house, a simple fence, two or three traffic cones, a big rounded bush"),
    ("Betongbil", "betongbil6",
     "a cement mixer truck at a construction site, its big round drum turning, seen from the side at a slight three-quarter angle.",
     "a truck with a cab with a big windshield and side window, a big tilted mixing drum with two or three wide spiral stripes, a chute at the back, three big wheels on each side",
     "the whole truck", "vehicle",
     "the half-built wall, the fence, traffic cones, the ground",
     "The stripes on the drum and along the cab",
     [BYGGE],
     "tires and dark parts",
     "side mirrors, headlights, a ladder up to the drum, a colored stripe along the cab",
     "a simple half-built wall, a simple fence, two or three traffic cones, a big rounded bush"),
    ("Kranbil", "kranbil4",
     "a mobile crane truck at a construction site, lifting a big wooden crate with its hook, seen from the side at a slight three-quarter angle.",
     "a long truck with a cab with a big windshield and side window, a solid telescopic crane boom (not a lattice), a hook on a short cable, outriggers resting on the ground, four big wheels on each side",
     "the whole truck and the whole crane boom", "vehicle",
     "the half-built house, the fence, traffic cones, the ground",
     "Stripes and warning patterns on the truck and the boom",
     [BYGGE, "The cable is a single bold line, the hook is one simple shape."],
     "tires, the outrigger feet and dark parts",
     "a beacon on the cab roof, side mirrors, headlights, steps, a colored stripe along the cab",
     "a simple half-built house with big window openings, a simple fence, two or three traffic cones, a big rounded bush"),
    # Ny skördetröska: fältet i skördetröskorna 2-5 blev pärlband och krumelurer
    # vid spårningen, så nu stora släta ytor i stället för vågor och strån
    ("Skördetröska", "skordetroska6",
     "a big green combine harvester harvesting a golden wheat field on a sunny summer day, seen from the side at a slight three-quarter angle. In front of the harvester the wheat is still standing, behind it the field is already cut.",
     "a combine harvester with a high cab with big windows on top of a big boxy body, a wide cutting header low in front, an unloading pipe folded along the side, big front wheels and smaller back wheels",
     "the whole combine harvester", "vehicle",
     "the edge of the standing wheat, the cut field, the hills, the tree, the barn",
     "Stripes on the harvester, the reel and the parts of the header",
     ["The standing wheat is ONE large smooth area of golden yellow with a smooth, gently wavy top edge, no lines inside it at all: no wave lines, no single straws, no ears of wheat, no rows.",
      "The cut field behind the harvester is ONE large smooth area of pale beige, no stubble, no lines, no rows, no tracks.",
      "No flying straw, chaff, dust or grain anywhere, nothing coming out of the harvester.",
      "The header is one wide smooth box with a big reel above it drawn as a simple cylinder with only four wide flat paddles, no teeth, no thin bars, no spikes, no auger spiral."],
     "tires and dark parts",
     "a beacon on the cab roof, side mirrors, two headlights, a ladder up to the cab, one wide colored stripe along the body",
     "soft rolling green hills, a red barn with a big door, one big rounded tree"),
    ("Polisbil", "polisbil6",
     "a police car driving along a street in a small town, seen from the side at a slight three-quarter angle.",
     "a station wagon police car with a big windshield and big side windows, a blue light bar on the roof, four wheels",
     "the whole car", "vehicle",
     "houses, trees, the street, the curb, road markings",
     "The blue and yellow checkered stripe along the side (a single row of large squares) and the stripes",
     [GATA],
     "tires and dark parts",
     "a roof light bar, side mirrors, headlights, door handles",
     "a few simple houses with big windows, a big rounded bush or tree near the car"),
    ("Ambulans", "ambulans6",
     "an ambulance driving along a street in a small town on its way to a small hospital, seen from the side at a slight three-quarter angle.",
     "a van-shaped ambulance with a cab with a big windshield and side window, a tall box body with a back door and one big side window, blue lights on the roof, four wheels",
     "the whole ambulance", "vehicle",
     "the hospital, houses, trees, the street, the curb, road markings",
     "The yellow and green checkered stripe along the side (a single row of large squares) and a simple cross symbol",
     [GATA],
     "tires and dark parts",
     "roof lights, side mirrors, headlights, door handles, a step at the back",
     "a small friendly hospital with big windows, a few simple houses, a big rounded bush or tree near the ambulance"),
    ("Sopbil", "sopbil8",
     "a garbage truck stopping by two wheelie bins on a street in a small town, seen from the side at a slight three-quarter angle.",
     "a garbage truck with a cab with a big windshield and side window, a big rounded container body, a lifting arm at the back, three big wheels on each side",
     "the whole truck and the bins", "vehicle",
     "houses, trees, the street, the curb, the bins, road markings",
     "Stripes and warning patterns on the truck",
     [GATA],
     "tires and dark parts",
     "a beacon on the cab roof, side mirrors, headlights, steps, a colored stripe along the body",
     "a few simple houses with big windows, a big rounded bush or tree near the truck"),
    ("Postbil", "postbil2",
     "a small yellow mail van driving along a street in a small town, with a few parcels visible through the back window, seen from the side at a slight three-quarter angle.",
     "a small delivery van with a big windshield, a side window, a sliding side door, four wheels",
     "the whole van", "vehicle",
     "houses, mailboxes, trees, the street, the curb, road markings",
     "A simple blue horn symbol (no letters) and the stripes on the van",
     [GATA],
     "tires and dark parts",
     "side mirrors, headlights, door handles, a colored stripe along the side",
     "a few simple houses with big windows, a mailbox on a post, a big rounded bush or tree near the van"),
    ("Buss", "buss4",
     "a city bus stopping at a bus stop on a street in a small town, seen from the side at a slight three-quarter angle.",
     "a long city bus with a big windshield, a row of four or five big side windows, two doors, four wheels",
     "the whole bus", "vehicle",
     "houses, the bus stop shelter, trees, the street, the curb, road markings",
     "Stripes along the bus",
     [GATA],
     "tires and dark parts",
     "side mirrors, headlights, a destination sign without text, a colored stripe along the side",
     "a few simple houses with big windows, a simple bus stop shelter, a big rounded bush or tree"),
    ("Tåg", "tag3",
     "a modern passenger train with a locomotive and two carriages, driving along the tracks through the countryside, seen from the side at a slight three-quarter angle.",
     "a streamlined train with a big front windshield, a row of big side windows on each carriage, doors, wheels hidden under a skirt",
     "the whole train", "vehicle",
     "hills, trees, the track, the embankment, the signal",
     "Stripes along the train",
     ["The track is two bold rails on one plain gravel embankment with only a few wide sleepers, no single stones."],
     "the wheels, the pantograph and dark parts",
     "headlights, a pantograph on the roof, door outlines, a colored stripe along the side",
     "soft rolling green hills, a simple railway signal, a big rounded bush or tree, a small red house far away"),
    ("Litet flygplan", "flygplan7",
     "a small propeller airplane standing on a grass airfield, ready for takeoff, seen from the side at a slight three-quarter angle.",
     "a small high-wing airplane with a big cockpit windshield and side window, a propeller with two blades, a tail fin, three wheels",
     "the whole airplane", "aircraft",
     "the hangar, the windsock, hills, trees, the grass field",
     "Stripes along the airplane",
     ["The grass airfield is one plain green area with one wide outlined runway strip, no single blades of grass."],
     "tires, the propeller and dark parts",
     "a colored stripe along the body, a colored tail, wheel covers, a door",
     "a simple hangar, a windsock on a pole, soft hills, a big rounded bush"),
    ("Helikopter", "helikopter4",
     "a rescue helicopter flying low over a green meadow by a lake, seen from the side at a slight three-quarter angle.",
     "a helicopter with a big rounded cockpit windshield, a side window and a sliding door, a main rotor with four wide blades, a tail boom with a small tail rotor, two landing skids",
     "the whole helicopter and its rotor", "aircraft",
     "the meadow, the lake, hills, trees",
     "Stripes along the helicopter",
     [VATTEN, "The rotor blades are solid shapes, no motion blur."],
     "the skids, rotor blades and dark parts",
     "a light under the body, a colored stripe along the body, a door handle",
     "soft rolling green hills, a small lake, a big rounded tree"),
    ("Raket", "raket4",
     "a big space rocket standing on its launch pad, ready to launch, seen at a slight three-quarter angle.",
     "a tall rocket with a pointed nose, a few round windows, two side boosters, big fins, standing on a simple launch tower",
     "the whole rocket and the launch tower", "rocket",
     "the launch tower, the ground, hills, the control building",
     "Stripes and bands on the rocket",
     ["A few big soft puffy clouds of smoke at the base, no flames, no sparks.",
      "The launch tower is a few solid beams, not a lattice."],
     "dark parts",
     "a few round windows, a hatch, colored bands around the body",
     "a simple control building, soft hills, a light blue sky with a pale moon"),
    ("Båt", "bat3",
     "a small motorboat with a cabin driving across a calm lake, seen from the side at a slight three-quarter angle.",
     "a small motorboat with a cabin with a big windshield and round side windows, a railing with only a few posts, a small flag",
     "the whole boat", "boat",
     "the water, the island, the jetty, hills, trees",
     "The stripe along the hull and a life ring",
     [VATTEN],
     "dark parts",
     "a life ring, a small flag, a colored stripe along the hull, a searchlight on the cabin roof",
     "a small island with a red cottage, a simple wooden jetty, soft hills, a big rounded tree"),
    ("Segelbåt", "segelbat3",
     "a sailboat sailing on a calm lake on a sunny day, seen from the side at a slight three-quarter angle.",
     "a sailboat with a big white main sail and a smaller front sail, a tall mast, a hull with a small cabin and round windows",
     "the whole boat and its sails", "boat",
     "the water, the island, hills, trees",
     "The stripe along the hull and the stripes on the sail",
     [VATTEN, "No thin ropes, at most one or two bold lines from the mast."],
     "dark parts",
     "a small flag on top of the mast, round windows, a colored stripe along the hull",
     "a small island with a red cottage and a lighthouse, soft hills, a big rounded tree"),
    ("Räddningsbåt", "raddningsbat2",
     "a rescue boat speeding across the sea near the coast, seen from the side at a slight three-quarter angle.",
     "a rescue boat with a cabin with a big windshield and side windows, a short mast with a blue light, a railing with only a few posts, a life ring",
     "the whole boat", "boat",
     "the water, the coast, the lighthouse, rocks",
     "The diagonal stripes on the hull and the life ring",
     [VATTEN],
     "dark parts",
     "a searchlight, a blue light on the mast, a life ring, a colored stripe along the hull",
     "a rocky coast with a few big smooth rocks, a red and white lighthouse, a small red cottage"),
    ("Lastfartyg", "lastfartyg2",
     "a cargo ship loaded with colorful containers sailing on the sea near a harbor, seen from the side at a slight three-quarter angle.",
     "a cargo ship with a long hull, a tall white bridge with a row of big windows at the back, two rows of large containers (no more than about ten containers), a funnel",
     "the whole ship", "ship",
     "the water, the harbor crane, the coast, hills",
     "Each container is a separate area with its own outline, with no ribs, and the stripe along the hull",
     [VATTEN],
     "dark parts",
     "a funnel with a colored band, an anchor, a colored stripe along the hull, a mast with a light",
     "a simple harbor crane made of solid beams, a coast with soft hills, a small lighthouse"),
]

STIL = [
    "Bold, clean, smooth outlines in near-black around EVERY shape, all the same thickness, in the foreground AND the background: {konturer}, clouds, the sun and each sun ray all get the same bold outline. No thin, light or colored lines anywhere.",
    "Every area must be fully closed by outlines, no gaps. {monster} are separate areas with their own outlines.",
]
STIL2 = [
    "Only the outlines are black. All fill colors are lighter: {mork} are medium gray, never black or dark navy.",
    "Mostly flat colors: each area has one main color with at most a very soft, light highlight. No strong gradients, no shading bands, no dark shadows along outlines, no shadows on the ground under the {fordon}.",
    "A bit more detail than a toy ({detaljer}), but no tiny enclosed areas. No grilles with many slots, no ribbed or accordion details, no panel lines, bolts or rivets, no tiny chain links.",
]
SIST = "No people, no animals, no text, no letters, no numbers, no logos, no frame, no signature, no paper texture."


def prompt(motiv, prop, hela, fordon, konturer, monster, extra, mork, detaljer, bakgrund):
    rader = ["A playful but slightly more realistic picture-book illustration for small children: " + motiv
             + " Realistic proportions: " + prop + ", but still friendly and softly rounded, not a toy."
             + " Landscape format 3:2, wider than tall, " + hela + " visible with free space all around.",
             "", "Style rules (important):"]
    regler = [s.format(konturer=konturer, monster=monster) for s in STIL] + extra
    regler += [s.format(mork=mork, fordon=fordon, detaljer=detaljer) for s in STIL2]
    regler += ["Background: " + bakgrund + ", light blue sky with a few big puffy clouds and a sun. "
               "Keep background objects large and simple, no other vehicles.", SIST]
    return "\n".join(rader + ["- " + r for r in regler])


INLEDNING = "Jag gör en målarboksapp för små barn. Bilderna görs om till målarbilder där barnet målar fram färgerna, så konturerna är det viktigaste."
REGLER = [
    "Skicka varje prompt **exakt som den står**, på engelska. Skriv inte om, korta inte och lägg inte till något.",
    "Liggande format **3:2** (till exempel 1536 × 1024). Aldrig kvadratiskt eller stående. Går formatet att välja i bildverktyget: välj 3:2 (eller 16:10/4:3 om 3:2 inte finns).",
    "En bild per prompt, i den ordning de står.",
    "Titta på varje bild innan du går vidare. Gör om den (högst två försök till) om den har text eller bokstäver, är kvadratisk eller stående, har människor eller djur, eller om fordonet inte syns helt.",
]


def markdown():
    d = ["# Bilder till målarboken Pensla", "", INLEDNING, "",
         "## Uppdrag", "",
         "Skapa alla bilderna nedan, **en i taget**, med bildgenereringen (Gemini). "
         "Spara varje bild i mappen `inkorg/` (skapa den om den saknas) med filnamnet som står "
         "i rubriken, till exempel `inkorg/traktor11.png`. Gå sedan direkt vidare till nästa "
         "bild utan att vänta på mig. Rör inga andra filer i projektet.", ""]
    d += ["- " + r for r in REGLER]
    d += ["- När alla är klara: skriv en lista över de sparade filerna och vilka som fick göras om.", ""]
    for i, (rubrik, fil, *resten) in enumerate(BILDER, 1):
        d += [f"## Bild {i}: {rubrik} → `inkorg/{fil}.png`", "", "```", prompt(*resten), "```", ""]
    return "\n".join(d)


def word():
    doc = Document()
    st = doc.styles['Normal']
    st.font.name = 'Calibri'
    st.font.size = Pt(11)
    doc.add_heading('Bilder till målarboken Pensla', 0)
    doc.add_paragraph(INLEDNING + " Skapa en bild för varje prompt i det här dokumentet.")
    doc.add_heading('Så här vill jag ha det', 1)
    for r in REGLER + ["Om du bara kan göra en bild åt gången: gör bild 1 och vänta tills jag skriver \"nästa\".",
                       "Ge varje bild filnamnet som står vid rubriken, till exempel traktor11.png."]:
        doc.add_paragraph(r.replace('**', ''), style='List Bullet')
    for i, (rubrik, fil, *resten) in enumerate(BILDER, 1):
        doc.add_heading(f'Bild {i}: {rubrik}  ({fil}.png)', 1)
        r = doc.add_paragraph().add_run(prompt(*resten))
        r.font.name = 'Consolas'
        r.font.size = Pt(9.5)
    for s in doc.sections:
        s.left_margin = s.right_margin = Cm(2)
    doc.save(UT + 'bilder.docx')


finns = {f.rsplit('.', 1)[0] for f in os.listdir(os.path.join(ROT, 'bilder', 'farg'))}
upptagna = [fil for _, fil, *_ in BILDER if fil in finns]
assert not upptagna, 'filnamnen finns redan: ' + ', '.join(upptagna)
with open(UT + 'bilder.md', 'w', encoding='utf-8', newline='\n') as f:
    f.write(markdown())
word()
print(len(BILDER), 'bilder')
